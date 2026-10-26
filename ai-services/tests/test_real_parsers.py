"""Tests validating binary PDF and DOCX parsing pipelines with layout structures."""

import io

from src.parsers.docx_parser import DOCXParser
from src.parsers.pdf_parser import PDFParser

try:
    import docx
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    from pypdf import PdfWriter
    HAS_PYPDF = True
except ImportError:
    HAS_PYPDF = False


def test_docx_table_and_paragraph_extraction():
    """Verify DOCX parser extracts both paragraphs and table columns."""
    if not HAS_DOCX:
        return

    doc = docx.Document()
    doc.add_heading("Senior Distributed Systems Engineer", level=1)
    doc.add_paragraph("Summary: Experienced engineer in cloud infrastructure.")

    # Add experience table (common layout format)
    table = doc.add_table(rows=2, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Role"
    hdr_cells[1].text = "Company"
    hdr_cells[2].text = "Dates"

    row_cells = table.rows[1].cells
    row_cells[0].text = "Staff Architect"
    row_cells[1].text = "Cloudscale Inc"
    row_cells[2].text = "2020 - Present"

    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)

    parser = DOCXParser()
    res = parser.parse_stream(stream)

    assert "Senior Distributed Systems Engineer" in res["text"]
    assert "Cloudscale Inc" in res["text"]
    assert "Staff Architect" in res["text"]
    assert res["char_count"] > 50


def test_pdf_extraction_fallback():
    """Verify PDF parser parses valid stream without crashing."""
    if not HAS_PYPDF:
        return

    # Create a minimal valid PDF stream
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    stream = io.BytesIO()
    writer.write(stream)
    stream.seek(0)

    parser = PDFParser()
    res = parser.parse_stream(stream)
    assert res["page_count"] == 1
